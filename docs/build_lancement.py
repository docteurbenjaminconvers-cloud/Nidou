"""Adds the 'Lancement 3 mois' tab: month-by-month cash forecast with growing sign-ups and the reverse trial."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = "/home/claude/nidou/docs/Nidou-previsionnel.xlsx"
wb = load_workbook(SRC)
if "Lancement 3 mois" in wb.sheetnames: del wb["Lancement 3 mois"]
ws = wb.create_sheet("Lancement 3 mois", 1)
F = lambda **k: Font(name="Arial", **k)
BLUE, GREEN = F(color="0000FF"), F(color="008000")
YELLOW, GREY = PatternFill("solid", fgColor="FFFF00"), PatternFill("solid", fgColor="EEEEEE")
EUR = '#,##0 €;(#,##0 €);-'; EUR2 = '#,##0.00 €;(#,##0.00 €);-'; INT = '#,##0;(#,##0);-'; PCT = '0.0%;(0.0%);-'
ws.column_dimensions["A"].width = 60
for c in "BCDE": ws.column_dimensions[c].width = 15
ws.column_dimensions["F"].width = 70

ws["A1"] = "Nidou — lancement : ce que coûtent les 3 premiers mois"; ws["A1"].font = F(bold=True, size=13)
ws["A2"] = "Cellules bleues sur fond jaune : modifiables. Hypothèses de coût unitaire reprises de l'onglet Hypothèses (vert)."; ws["A2"].font = F(italic=True, size=9)

# ---- inputs
r = 4
ws.cell(r, 1, "HYPOTHÈSES DE LANCEMENT").font = F(bold=True); ws.cell(r, 1).fill = GREY; r += 1
inputs = [
    ("Inscriptions nouvelles par mois", [500, 1500, 4000], INT, "Liste d'attente, bouche-à-oreille, premiers partages ; sans publicité payante"),
    ("Part des inscrits qui convertissent en payant à la fin de l'essai", [0.08, 0.08, 0.08], PCT, "Essai inversé 15 jours sans carte, puis gratuit nu (sans IA ni banque) : 5 à 12 % selon les produits ; 8 % en hypothèse centrale"),
    ("Churn mensuel des payants", [0.03, 0.03, 0.03], PCT, ""),
    ("Part des inscrits en essai qui connectent leur banque", [0.5, 0.5, 0.5], PCT, "Pendant l'essai, Bridge est ouvert : c'est le principal coût de l'essai"),
    ("Documents lus par l'IA pendant l'essai, par inscrit", [8, 8, 8], INT, "Modèle rapide (0,5 c) ; 15 % repassent en modèle précis (2 c)"),
    ("Documents lus par l'IA par compte gratuit et par mois (après essai)", [0, 0, 0], INT, "Gratuit nu : lecture locale sans IA, pas d'import de relevés ni de banque (NIDOU_IA_LECTURES_GRATUIT=0)"),
]
inp = {}
for label, vals, fmt, note in inputs:
    ws.cell(r, 1, label).font = F()
    for j, v in enumerate(vals):
        c = ws.cell(r, 2 + j, v); c.font = BLUE; c.fill = YELLOW; c.number_format = fmt
    ws.cell(r, 6, note).font = F(size=9, color="555555")
    inp[label] = r; r += 1
fixed = [
    ("Hébergement OVH", [60, 60, 120], EUR, "VPS puis une instance de plus au mois 3"),
    ("Juridique, comptable, assurance (dont dépôt INPI 190 € au mois 1)", [900, 500, 500], EUR, "Expert-comptable ~250 €/mois, RC pro, CGV/RGPD relus par un avocat (~1 500 € à prévoir une fois, non inclus)"),
    ("Outils (domaine, Brevo, monitoring, Stripe hors commissions)", [100, 100, 150], EUR, ""),
    ("Marketing (contenu, visuels, petits partenariats)", [300, 500, 1000], EUR, "Pas de publicité payante avant une conversion mesurée"),
    ("Équipe (hors fondateur)", [0, 0, 0], EUR, "Le fondateur et Claude ; une première embauche n'est pas nécessaire avant ~5 000 payants"),
]
fx = {}
for label, vals, fmt, note in fixed:
    ws.cell(r, 1, label).font = F()
    for j, v in enumerate(vals):
        c = ws.cell(r, 2 + j, v); c.font = BLUE; c.fill = YELLOW; c.number_format = fmt
    ws.cell(r, 6, note).font = F(size=9, color="555555")
    fx[label] = r; r += 1

# references to the Hypothèses tab (unit costs & price)
hy = wb["Hypothèses"]; HL = {hy.cell(i, 1).value: i for i in range(1, hy.max_row + 1)}
H = lambda label: f"Hypothèses!$B${HL[label]}"
r += 1
ws.cell(r, 1, "COÛTS UNITAIRES (repris de l'onglet Hypothèses)").font = F(bold=True); ws.cell(r, 1).fill = GREY; r += 1
units = [
    ("Coût d'une lecture IA pendant l'essai (€)", "=0.85*0.005+0.15*0.02", EUR2, "85 % modèle rapide à 0,5 c, 15 % modèle précis à 2 c"),
    ("Bridge par compte connecté et par mois (€)", f"={H('Agrégation bancaire Bridge, par compte connecté (€)')}", EUR2, ""),
    ("Comptes connectés par foyer", f"={H('Nombre de comptes connectés par foyer connecté')}", "0.0", ""),
    ("Stockage et push par compte et par mois (€)", f"={H('Stockage et bande passante, tous foyers (€)')}", EUR2, "Push gratuit ; pas d'email sur le gratuit"),
    ("Revenu net par payant et par mois (€)", "=Prévisionnel!$B$14/Prévisionnel!$B$4", EUR2, "Après TVA et commissions Stripe"),
    ("Coût variable d'un payant par mois (€)", "=-Prévisionnel!$B$21/Prévisionnel!$B$4", EUR2, "IA, Bridge, stockage, emails"),
]
un = {}
for label, f, fmt, note in units:
    ws.cell(r, 1, label).font = F()
    c = ws.cell(r, 2, f); c.font = GREEN; c.number_format = fmt
    ws.cell(r, 6, note).font = F(size=9, color="555555")
    un[label] = r; r += 1

# ---- monthly model
r += 1
ws.cell(r, 1, "MOIS").font = F(bold=True); ws.cell(r, 1).fill = GREY
for j, m in enumerate(["Mois 1", "Mois 2", "Mois 3"]):
    c = ws.cell(r, 2 + j, m); c.font = F(bold=True); c.fill = GREY
ws.cell(r, 5, "Total 3 mois").font = F(bold=True); ws.cell(r, 5).fill = GREY
head = r; r += 1
L = lambda j: get_column_letter(2 + j)
I = lambda label: inp[label]
def row(label, fn, fmt=EUR, bold=False, total=True, font=F()):
    global r
    ws.cell(r, 1, label).font = F(bold=bold)
    for j in range(3):
        c = ws.cell(r, 2 + j, fn(j)); c.number_format = fmt; c.font = F(bold=bold, color=font.color)
    if total:
        c = ws.cell(r, 5, f"=SUM(B{r}:D{r})"); c.number_format = fmt; c.font = F(bold=True)
    r += 1
    return r - 1

K_CONV = I("Part des inscrits qui convertissent en payant à la fin de l'essai")
K_CHURN = I("Churn mensuel des payants")
K_BANK = I("Part des inscrits en essai qui connectent leur banque")
K_DOCS_ESSAI = I("Documents lus par l'IA pendant l'essai, par inscrit")
K_DOCS_FREE = I("Documents lus par l'IA par compte gratuit et par mois (après essai)")
U_IA = un["Coût d'une lecture IA pendant l'essai (€)"]
U_BR = un["Bridge par compte connecté et par mois (€)"]
U_CPT = un["Comptes connectés par foyer"]
U_STO = un["Stockage et push par compte et par mois (€)"]
U_REV = un["Revenu net par payant et par mois (€)"]
U_VAR = un["Coût variable d'un payant par mois (€)"]
ROW_INSC = row("Inscriptions du mois", lambda j: f"={L(j)}{I('Inscriptions nouvelles par mois')}", INT)
ROW_ESSAI = row("Comptes en essai (15 jours, plan Famille complet)", lambda j: f"={L(j)}{ROW_INSC}", INT, total=False)
ROW_CONV = row("Nouveaux payants (conversion en fin d'essai)", lambda j: f"=ROUND({L(j)}{ROW_INSC}*{L(j)}{K_CONV},0)", INT)
# simpler: compute payants row with explicit references
ws.cell(r, 1, "Payants en fin de mois").font = F(bold=True)
ROW_PAY = r
for j in range(3):
    f = f"={L(j)}{ROW_CONV}" if j == 0 else f"=ROUND({L(j-1)}{ROW_PAY}*(1-{L(j)}{K_CHURN}),0)+{L(j)}{ROW_CONV}"
    c = ws.cell(r, 2 + j, f); c.number_format = INT; c.font = F(bold=True)
r += 1
ws.cell(r, 1, "Comptes gratuits actifs en fin de mois (après essai, non convertis)").font = F()
ROW_FREE = r
for j in range(3):
    f = f"={L(j)}{ROW_INSC}-{L(j)}{ROW_CONV}" if j == 0 else f"={L(j-1)}{ROW_FREE}+{L(j)}{ROW_INSC}-{L(j)}{ROW_CONV}"
    c = ws.cell(r, 2 + j, f); c.number_format = INT
r += 2

ws.cell(r, 1, "REVENUS").font = F(bold=True); ws.cell(r, 1).fill = GREY; r += 1
ROW_REV = row("Revenu net encaissé (payants moyens du mois × revenu net)", lambda j: f"=({L(j)}{ROW_PAY}-{L(j)}{ROW_CONV}/2)*$B${U_REV}", bold=True)
r += 1
ws.cell(r, 1, "COÛTS VARIABLES").font = F(bold=True); ws.cell(r, 1).fill = GREY; r += 1
ROW_C_IA_ESSAI = row("IA pendant l'essai", lambda j: f"=-{L(j)}{ROW_ESSAI}*{L(j)}{K_DOCS_ESSAI}*$B${U_IA}")
ROW_C_BR_ESSAI = row("Bridge pendant l'essai (½ mois, part connectée)", lambda j: f"=-{L(j)}{ROW_ESSAI}*{L(j)}{K_BANK}*$B${U_CPT}*$B${U_BR}*0.5")
ROW_C_FREE = row("Comptes gratuits (stockage + push ; IA à zéro)", lambda j: f"=-{L(j)}{ROW_FREE}*({L(j)}{K_DOCS_FREE}*0.005+$B${U_STO})")
ROW_C_PAY = row("Payants (IA, Bridge, stockage, emails)", lambda j: f"=-({L(j)}{ROW_PAY}-{L(j)}{ROW_CONV}/2)*$B${U_VAR}")
ROW_C_VAR = row("Total coûts variables", lambda j: f"=SUM({L(j)}{ROW_C_IA_ESSAI}:{L(j)}{ROW_C_PAY})", bold=True)
r += 1
ws.cell(r, 1, "COÛTS FIXES").font = F(bold=True); ws.cell(r, 1).fill = GREY; r += 1
first = r
for label in fx:
    row(label, lambda j, fr=fx[label]: f"=-{L(j)}{fr}", font=GREEN)
ROW_C_FIX = row("Total coûts fixes", lambda j: f"=SUM({L(j)}{first}:{L(j)}{r-1})", bold=True)
r += 1
ROW_RES = row("RÉSULTAT DU MOIS", lambda j: f"={L(j)}{ROW_REV}+{L(j)}{ROW_C_VAR}+{L(j)}{ROW_C_FIX}", bold=True)
for j in range(4): ws.cell(ROW_RES, 2 + j).fill = YELLOW
ws.cell(r, 1, "Trésorerie cumulée (besoin de financement si négatif)").font = F(bold=True)
ROW_CUM = r
for j in range(3):
    f = f"={L(j)}{ROW_RES}" if j == 0 else f"={L(j-1)}{ROW_CUM}+{L(j)}{ROW_RES}"
    c = ws.cell(r, 2 + j, f); c.number_format = EUR; c.font = F(bold=True)
r += 2
ws.cell(r, 1, "REPÈRES").font = F(bold=True); ws.cell(r, 1).fill = GREY; r += 1
row("Coût d'un inscrit en essai (IA + Bridge), €", lambda j: f"=IF({L(j)}{ROW_INSC}=0,0,-({L(j)}{ROW_C_IA_ESSAI}+{L(j)}{ROW_C_BR_ESSAI})/{L(j)}{ROW_INSC})", EUR2, total=False)
row("Coût d'un compte gratuit par mois, €", lambda j: f"=IF({L(j)}{ROW_FREE}=0,0,-{L(j)}{ROW_C_FREE}/{L(j)}{ROW_FREE})", EUR2, total=False)
row("Coût total des inscrits ÷ nouveaux payants (coût d'acquisition réel), €", lambda j: f"=IF({L(j)}{ROW_CONV}=0,0,-({L(j)}{ROW_C_IA_ESSAI}+{L(j)}{ROW_C_BR_ESSAI})/{L(j)}{ROW_CONV})", EUR2, total=False)
row("Chiffre d'affaires TTC facturé dans le mois (indicatif)", lambda j: f"=({L(j)}{ROW_PAY}-{L(j)}{ROW_CONV}/2)*Prévisionnel!$B$9")
r += 1
ws.cell(r, 1, "Lecture : l'essai inversé coûte ≈ 0,40 € par inscrit (surtout Bridge, ouvert 15 jours). À 8 % de conversion, cela revient à ≈ 5 € par payant gagné, pour un abonné qui rapporte ≈ 7 € nets dès son premier mois. Le gratuit nu ne coûte que le stockage et le push : ≈ 1 centime par compte et par mois, zéro IA. Les coûts fixes du lancement (comptable, outils, hébergement) pèsent plus que les utilisateurs.").font = F(size=9, italic=True)
ws.freeze_panes = "B5"
wb.save(SRC)
print("ok")
